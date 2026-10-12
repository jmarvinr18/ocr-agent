import io

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Picture


def _md_cell(text):
    return text.strip().replace("|", "\\|").replace("\n", "<br>")


def _md_table(rows):
    """Render a list of rows (lists of cell text) as a Markdown table."""
    rows = [[_md_cell(cell) for cell in row] for row in rows if row]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]

    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join(lines)


# --- .docx -----------------------------------------------------------------


def _docx_paragraph(paragraph, document, describe_image):
    blocks = []
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style is not None else ""

    if text:
        if style == "Title":
            blocks.append(f"# {text}")
        elif style.startswith("Heading"):
            level = style.split()[-1]
            level = int(level) if level.isdigit() else 1
            blocks.append(f"{'#' * min(level, 6)} {text}")
        elif style.startswith("List Number"):
            blocks.append(f"1. {text}")
        elif "List" in style:
            blocks.append(f"- {text}")
        else:
            blocks.append(text)

    # Inline pictures are referenced from the paragraph by relationship id.
    for blip in paragraph._p.iter(qn("a:blip")):
        part = document.part.related_parts.get(blip.get(qn("r:embed")))
        if part is not None:
            blocks.append(describe_image(part.blob, part.content_type))

    return "\n\n".join(b for b in blocks if b)


def docx_to_markdown(data, describe_image):
    """Convert a .docx to Markdown, keeping paragraphs and tables in document order."""
    document = docx.Document(io.BytesIO(data))
    blocks = []

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(_docx_paragraph(Paragraph(child, document), document, describe_image))
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            blocks.append(_md_table([[cell.text for cell in row.cells] for row in table.rows]))

    return "\n\n".join(b for b in blocks if b)


# --- .pptx -----------------------------------------------------------------


def _reading_order(shapes):
    return sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))


def _pptx_shape(shape, describe_image):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [b for s in _reading_order(shape.shapes) for b in _pptx_shape(s, describe_image)]

    if shape.has_table:
        return [_md_table([[cell.text for cell in row.cells] for row in shape.table.rows])]

    if isinstance(shape, Picture):
        try:
            image = shape.image
        except ValueError:  # linked, not embedded
            return []
        return [describe_image(image.blob, image.content_type)]

    if shape.has_text_frame:
        lines = [
            f"{'  ' * p.level}- {p.text.strip()}"
            for p in shape.text_frame.paragraphs
            if p.text.strip()
        ]
        return ["\n".join(lines)]

    return []


def pptx_to_markdown(data, describe_image):
    """Convert a .pptx to Markdown, one section per slide."""
    presentation = Presentation(io.BytesIO(data))
    slides = []

    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        blocks = [f"## Slide {number}" + (f": {title}" if title else "")]

        for shape in _reading_order(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            blocks.extend(_pptx_shape(shape, describe_image))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append("> **Notes:** " + notes.replace("\n", "\n> "))

        slides.append("\n\n".join(b for b in blocks if b))

    return "\n\n".join(slides)
